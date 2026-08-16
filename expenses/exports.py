import csv
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND = colors.HexColor("#4f46e5")
MUTED = colors.HexColor("#64748b")
BORDER = colors.HexColor("#e2e8f0")
STRIPE = colors.HexColor("#f8fafc")


def csv_response(filename, headers, rows):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow(headers)
    writer.writerows(rows)
    return response


def xlsx_response(filename, headers, rows, sheet_title="Data"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31] or "Data"

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        sheet.append(row)

    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells), default=0)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 42)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def pdf_summary_response(
    filename, title, subtitle, total_label, total_amount, count, breakdown_title, breakdown_rows,
    breakdown_columns=("Category", "Amount", "% of total"),
):
    """A one-page PDF summary: headline total plus a breakdown table -
    a report to read, not a raw dump of every row."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("TitleBrand", parent=styles["Title"], textColor=BRAND, spaceAfter=2)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], textColor=MUTED, fontSize=10)
    label_style = ParagraphStyle("Label", parent=styles["Normal"], textColor=MUTED, fontSize=10, spaceBefore=16)
    headline_style = ParagraphStyle("Headline", parent=styles["Title"], fontSize=30, leading=34, spaceBefore=2, spaceAfter=2)
    heading2_style = ParagraphStyle("H2", parent=styles["Heading2"], textColor=BRAND, spaceBefore=8)

    elements = [
        Paragraph(title, title_style),
        Paragraph(subtitle, meta_style),
        Spacer(1, 4),
        Paragraph(total_label, label_style),
        Paragraph(f"{total_amount:.2f}", headline_style),
        Paragraph(f"{count} transaction{'s' if count != 1 else ''}", meta_style),
    ]

    if breakdown_rows:
        elements.append(Paragraph(breakdown_title, heading2_style))
        elements.append(Spacer(1, 8))
        table_data = [list(breakdown_columns)] + breakdown_rows
        table = Table(table_data, colWidths=[3 * inch, 1.7 * inch, 1.5 * inch])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), BRAND),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                    ("TOPPADDING", (0, 0), (-1, 0), 8),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, STRIPE]),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.5, BORDER),
                    ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
