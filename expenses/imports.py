"""Bulk expense import from an Excel (.xlsx) or CSV upload.

The expected columns are the same ones the expense export writes, so an
exported file can be edited and re-imported as-is. Import is
all-or-nothing: every row is validated first, and if any row has a
problem nothing is saved and the errors are reported back by row number.
"""
import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO, TextIOWrapper

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import PAYMENT_TYPE_CHOICES, Category, Expense
from .utils import CATEGORY_COLUMNS, CATEGORY_PATH_SEPARATOR, category_path_cells, category_path_names

IMPORT_HEADERS = ["Date", *CATEGORY_COLUMNS, "Amount", "Payment Type", "Payment Source", "Location", "Note"]
REQUIRED_HEADERS = {"date", "category", "amount"}
MAX_ROWS = 5000
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

# Day-first for ambiguous dd/mm vs mm/dd strings. Excel date cells arrive
# as real datetimes and don't go through these at all.
DATE_FORMATS = [
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y %H:%M", "%d-%m-%Y",
]

_PAYMENT_LOOKUP = {}
for _value, _label in PAYMENT_TYPE_CHOICES:
    _PAYMENT_LOOKUP[_value.lower()] = _value
    _PAYMENT_LOOKUP[_label.lower()] = _value


class ImportFileError(Exception):
    """The upload as a whole can't be read (wrong type, no header row, ...)."""


@dataclass
class RowError:
    row: int
    message: str


@dataclass
class ImportResult:
    created: int = 0
    categories_created: int = 0
    errors: list = None


def _read_rows(upload):
    """Yield (row_number, [cell values]) for every row, header included."""
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except Exception:
            raise ImportFileError("That file couldn't be opened as an Excel workbook.")
        sheet = workbook.worksheets[0]
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            yield number, list(values)
    elif name.endswith(".csv"):
        try:
            reader = csv.reader(TextIOWrapper(upload.file, encoding="utf-8-sig"))
            for number, values in enumerate(reader, start=1):
                yield number, values
        except UnicodeDecodeError:
            raise ImportFileError("That CSV isn't UTF-8 encoded. Re-save it as “CSV UTF-8” and try again.")
    else:
        raise ImportFileError("Upload an Excel (.xlsx) or CSV file.")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _parse_date(value):
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time())
    else:
        text = _text(value)
        for fmt in DATE_FORMATS:
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"“{text}” isn't a date. Use YYYY-MM-DD or DD/MM/YYYY (optionally with HH:MM).")
    return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed


def _parse_amount(value):
    text = _text(value).replace(",", "")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"“{_text(value)}” isn't a number.")
    if amount <= 0:
        raise ValueError("Amount must be greater than zero.")
    return amount.quantize(Decimal("0.01"))


def _parse_payment_type(value):
    text = _text(value).lower()
    if not text:
        return "cash"
    if text not in _PAYMENT_LOOKUP:
        allowed = ", ".join(label for _, label in PAYMENT_TYPE_CHOICES)
        raise ValueError(f"Unknown payment type “{_text(value)}”. Use one of: {allowed}.")
    return _PAYMENT_LOOKUP[text]


def _category_parts(cells):
    """Root-to-leaf names from the Category / Subcategory 1-5 cells.

    Subcategories are optional but must be filled left to right. A cell
    may itself hold a "A > B" path (that's how export writes categories
    nested deeper than the columns allow); it's split into its levels."""
    texts = [_text(c) for c in cells]
    if not texts[0]:
        raise ValueError("Category is required.")
    filled = [t for t in texts if t]
    last_filled = max(i for i, t in enumerate(texts) if t)
    if len(filled) != last_filled + 1:
        gap = next(i for i, t in enumerate(texts) if not t)
        raise ValueError(
            f"{CATEGORY_COLUMNS[gap]} is empty but a later subcategory column is filled. Fill them left to right."
        )
    parts = [p.strip() for t in filled for p in t.split(CATEGORY_PATH_SEPARATOR.strip())]
    if not all(parts):
        raise ValueError("Category has an empty level in it.")
    return parts


class _CategoryResolver:
    """Resolves a root-to-leaf list of names to one of the user's
    categories.

    A multi-level path is matched level by level from the root. A single
    name (only the Category column filled) means that root category; if
    there's no such root, it may also name a subcategory anywhere in the
    tree, as long as only one category has that name. With
    create_missing, unknown levels are created on the fly."""

    def __init__(self, user, create_missing):
        self.user = user
        self.create_missing = create_missing
        self.created = 0
        self.categories = list(Category.objects.filter(user=user))
        self.by_id = {c.id: c for c in self.categories}

    def _children(self, parent_id):
        return [c for c in self.categories if c.parent_id == parent_id]

    def path_names(self, category):
        return category_path_names(category.id, self.by_id)

    def _path(self, category):
        return CATEGORY_PATH_SEPARATOR.join(self.path_names(category))

    def resolve(self, parts):
        if len(parts) == 1:
            name = parts[0].lower()
            root = next((c for c in self._children(None) if c.name.lower() == name), None)
            if root is not None:
                return root
            matches = [c for c in self.categories if c.name.lower() == name]
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                options = "; ".join(sorted(self._path(c) for c in matches))
                raise ValueError(
                    f"“{parts[0]}” matches several subcategories ({options}). "
                    "Fill in the Subcategory columns to say which one."
                )

        parent = None
        for depth, name in enumerate(parts):
            parent_id = parent.id if parent else None
            match = next((c for c in self._children(parent_id) if c.name.lower() == name.lower()), None)
            if match is None:
                if not self.create_missing:
                    raise ValueError(
                        f"Category “{CATEGORY_PATH_SEPARATOR.join(parts[:depth + 1])}” doesn't exist. "
                        "Create it first, or tick “Create categories that don't exist yet”."
                    )
                match = Category.objects.create(user=self.user, parent=parent, name=name)
                self.categories.append(match)
                self.by_id[match.id] = match
                self.created += 1
            parent = match
        return parent


def import_expenses(upload, user, create_missing=False):
    """Validate every row of `upload` and, if all are valid, create the
    expenses (and any missing categories) in one transaction. Raises
    ImportFileError for an unreadable file; per-row problems come back in
    ImportResult.errors with nothing saved."""
    if upload.size > MAX_UPLOAD_BYTES:
        raise ImportFileError("That file is larger than 5 MB.")

    rows = _read_rows(upload)
    header_number, header = next(
        ((n, values) for n, values in rows if any(_text(v) for v in values)), (None, None)
    )
    if header is None:
        raise ImportFileError("That file is empty.")
    columns = {_text(name).lower(): index for index, name in enumerate(header) if _text(name)}
    missing = [h for h in IMPORT_HEADERS if h.lower() in REQUIRED_HEADERS and h.lower() not in columns]
    if missing:
        raise ImportFileError(
            f"Row {header_number} should be the header row, but it's missing: {', '.join(missing)}. "
            "Download the template to see the expected columns."
        )

    def cell(values, header_name):
        index = columns.get(header_name.lower())
        return values[index] if index is not None and index < len(values) else None

    errors, expenses = [], []
    with transaction.atomic():
        resolver = _CategoryResolver(user, create_missing)
        for number, values in rows:
            if not any(_text(v) for v in values):
                continue
            if len(expenses) + len(errors) >= MAX_ROWS:
                raise ImportFileError(f"That file has more than {MAX_ROWS} rows. Split it into smaller files.")
            try:
                expense = Expense(
                    user=user,
                    date=_parse_date(cell(values, "Date")),
                    category=resolver.resolve(_category_parts([cell(values, h) for h in CATEGORY_COLUMNS])),
                    amount=_parse_amount(cell(values, "Amount")),
                    payment_type=_parse_payment_type(cell(values, "Payment Type")),
                    payment_source=_text(cell(values, "Payment Source")),
                    location=_text(cell(values, "Location")),
                    note=_text(cell(values, "Note")),
                )
                expense.full_clean(exclude=["user", "category"])
            except ValueError as exc:
                errors.append(RowError(number, str(exc)))
                continue
            except ValidationError as exc:
                for field, messages in exc.message_dict.items():
                    label = field.replace("_", " ").capitalize()
                    errors.extend(RowError(number, f"{label}: {m}") for m in messages)
                continue
            expenses.append(expense)

        if errors:
            transaction.set_rollback(True)
            return ImportResult(errors=errors)
        if not expenses:
            raise ImportFileError("That file has a header row but no expenses under it.")

        Expense.objects.bulk_create(expenses)
    return ImportResult(created=len(expenses), categories_created=resolver.created, errors=[])


def template_workbook_bytes(user):
    """A ready-to-fill .xlsx: the import headers (with a hover note on
    each explaining what goes there), a payment-type dropdown, and a
    second sheet listing the user's categories by full path. No example
    row, so a forgotten sample can't be imported as a real expense."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    sheet.append(IMPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    notes = [
        "Required. An Excel date, or text like 2026-10-18, 2026-10-18 14:30 or 18/10/2026.",
        "Required. The top-level category, e.g. Food. See the “Your categories” sheet.",
        *(
            f"Optional. Level {level} under Category, e.g. "
            + ("Dining Out" if level == 1 else "Coffee" if level == 2 else "a deeper subcategory")
            + ". Fill subcategory columns left to right."
            for level in range(1, len(CATEGORY_COLUMNS))
        ),
        "Required. Greater than zero.",
        "Optional. Pick from the dropdown. Blank means Cash.",
        "Optional. e.g. HDFC Bank",
        "Optional.",
        "Optional.",
    ]
    for cell, note in zip(sheet[1], notes):
        cell.comment = Comment(note, "Finlytics")

    payment_column = get_column_letter(IMPORT_HEADERS.index("Payment Type") + 1)
    payment_labels = ",".join(label for _, label in PAYMENT_TYPE_CHOICES)
    validation = DataValidation(type="list", formula1=f'"{payment_labels}"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"{payment_column}2:{payment_column}{MAX_ROWS + 1}")

    widths = {"Date": 18, "Amount": 12, "Payment Type": 15, "Payment Source": 20, "Location": 20, "Note": 32}
    for index, header in enumerate(IMPORT_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 16)
    sheet.freeze_panes = "A2"

    # Every existing category as a row of cells laid out exactly like the
    # import columns, so a row can be copied straight across.
    resolver = _CategoryResolver(user, create_missing=False)
    paths = sorted((resolver.path_names(c) for c in resolver.categories), key=lambda names: [n.lower() for n in names])
    reference = workbook.create_sheet("Your categories")
    reference.append(CATEGORY_COLUMNS)
    for cell in reference[1]:
        cell.font = Font(bold=True)
    for names in paths:
        reference.append(category_path_cells(names))
    for index in range(1, len(CATEGORY_COLUMNS) + 1):
        reference.column_dimensions[get_column_letter(index)].width = 18

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
