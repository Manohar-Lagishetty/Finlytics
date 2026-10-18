from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from .models import Category, Expense


class BulkDeleteExpensesTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(email="me@example.com", password="pw", is_active=True)
        self.other = User.objects.create_user(email="other@example.com", password="pw")
        self.category = Category.objects.create(user=self.user, name="Food")
        other_category = Category.objects.create(user=self.other, name="Food")
        self.mine = [
            Expense.objects.create(user=self.user, category=self.category, amount=Decimal("10.00"))
            for _ in range(3)
        ]
        self.theirs = Expense.objects.create(user=self.other, category=other_category, amount=Decimal("5.00"))
        self.client.force_login(self.user)
        self.url = reverse("expenses:bulk_delete")

    def test_deletes_selected_expenses_only(self):
        ids = [self.mine[0].pk, self.mine[1].pk]
        response = self.client.post(self.url, {"ids": ids, "next": "/expenses/?page=2"})

        self.assertRedirects(response, "/expenses/?page=2", fetch_redirect_response=False)
        self.assertEqual(list(Expense.objects.filter(user=self.user)), [self.mine[2]])

    def test_ignores_other_users_expenses(self):
        self.client.post(self.url, {"ids": [self.mine[0].pk, self.theirs.pk]})

        self.assertTrue(Expense.objects.filter(pk=self.theirs.pk).exists())
        self.assertFalse(Expense.objects.filter(pk=self.mine[0].pk).exists())

    def test_ignores_non_numeric_ids_and_empty_selection(self):
        response = self.client.post(self.url, {"ids": ["abc", ""]})

        self.assertRedirects(response, reverse("expenses:list"), fetch_redirect_response=False)
        self.assertEqual(Expense.objects.count(), 4)

    def test_rejects_offsite_next(self):
        response = self.client.post(self.url, {"ids": [self.mine[0].pk], "next": "https://evil.example"})

        self.assertRedirects(response, reverse("expenses:list"), fetch_redirect_response=False)

    def test_requires_post(self):
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 405)

    def test_select_all_deletes_every_matching_expense(self):
        cheap = Expense.objects.create(user=self.user, category=self.category, amount=Decimal("1.00"))
        response = self.client.post(
            self.url, {"select_all": "1", "filters": "amount_min=5&sort=-date&page=2", "next": "/expenses/"}
        )

        self.assertRedirects(response, "/expenses/", fetch_redirect_response=False)
        self.assertEqual(list(Expense.objects.filter(user=self.user)), [cheap])
        self.assertTrue(Expense.objects.filter(pk=self.theirs.pk).exists())

    def test_select_all_without_filters_deletes_all_of_my_expenses(self):
        self.client.post(self.url, {"select_all": "1", "filters": ""})

        self.assertFalse(Expense.objects.filter(user=self.user).exists())
        self.assertTrue(Expense.objects.filter(pk=self.theirs.pk).exists())

    def test_select_all_with_invalid_filters_deletes_nothing(self):
        self.client.post(self.url, {"select_all": "1", "filters": "amount_min=notanumber"})

        self.assertEqual(Expense.objects.filter(user=self.user).count(), 3)


def _xlsx_upload(rows, name="expenses.xlsx"):
    from openpyxl import Workbook

    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


HEADERS = [
    "Date", "Category", "Subcategory 1", "Subcategory 2", "Subcategory 3", "Subcategory 4", "Subcategory 5",
    "Amount", "Payment Type", "Payment Source", "Location", "Note",
]


def _row(date, categories, amount, payment_type="", payment_source="", location="", note=""):
    """One sheet row: `categories` fills Category then Subcategory 1-5."""
    cells = list(categories) + [None] * (6 - len(categories))
    return [date, *cells, amount, payment_type, payment_source, location, note]


class ImportExpensesTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(email="me@example.com", password="pw", is_active=True)
        self.food = Category.objects.create(user=self.user, name="Food")
        self.dining = Category.objects.create(user=self.user, name="Dining Out", parent=self.food)
        self.coffee = Category.objects.create(user=self.user, name="Coffee", parent=self.dining)
        self.client.force_login(self.user)
        self.url = reverse("expenses:import")

    def post(self, upload, **extra):
        return self.client.post(self.url, {"file": upload, **extra})

    def test_imports_valid_rows(self):
        upload = _xlsx_upload([
            HEADERS,
            _row(datetime(2026, 10, 1, 9, 30), ["Food", "Dining Out", "Coffee"], 120, "UPI", "HDFC", "Cafe", "latte"),
            _row("2026-10-02", ["food", "dining out"], "80.5"),
            [None] * len(HEADERS),  # blank rows are skipped
            _row("03/10/2026", ["Food"], 1500, "bank transfer"),
        ])
        response = self.post(upload)

        self.assertRedirects(response, reverse("expenses:list"), fetch_redirect_response=False)
        expenses = list(Expense.objects.filter(user=self.user).order_by("date"))
        self.assertEqual([e.amount for e in expenses], [Decimal("120.00"), Decimal("80.50"), Decimal("1500.00")])
        self.assertEqual([e.category for e in expenses], [self.coffee, self.dining, self.food])
        self.assertEqual([e.payment_type for e in expenses], ["upi", "cash", "bank_transfer"])
        self.assertEqual(expenses[2].date.day, 3)  # 03/10 is read day-first

    def test_bare_subcategory_name_in_category_column_still_resolves(self):
        self.post(_xlsx_upload([HEADERS, _row("2026-10-01", ["Coffee"], 10)]))

        self.assertEqual(Expense.objects.get().category, self.coffee)

    def test_round_trips_the_excel_export(self):
        Expense.objects.create(user=self.user, category=self.coffee, amount=Decimal("42.00"), note="x")
        export = self.client.get(reverse("expenses:export_xlsx"))
        Expense.objects.all().delete()

        self.post(SimpleUploadedFile("expenses.xlsx", export.content))

        expense = Expense.objects.get(user=self.user)
        self.assertEqual((expense.category, expense.amount, expense.note), (self.coffee, Decimal("42.00"), "x"))

    def test_export_writes_one_level_per_column_and_folds_deep_paths(self):
        from openpyxl import load_workbook

        parent = self.coffee
        for name in ["L4", "L5", "L6", "L7"]:  # 7 levels deep: one more than the 6 columns
            parent = Category.objects.create(user=self.user, name=name, parent=parent)
        Expense.objects.create(user=self.user, category=parent, amount=Decimal("1.00"))
        export = self.client.get(reverse("expenses:export_xlsx"))

        sheet = load_workbook(BytesIO(export.content)).active
        self.assertEqual([c.value for c in sheet[1]], HEADERS)
        self.assertEqual([c.value for c in sheet[2]][1:7], ["Food", "Dining Out", "Coffee", "L4", "L5", "L6 > L7"])

        Expense.objects.all().delete()
        self.post(SimpleUploadedFile("expenses.xlsx", export.content))
        self.assertEqual(Expense.objects.get().category, parent)

    def test_imports_csv_with_only_required_columns(self):
        csv_bytes = "Date,Category,Amount\n2026-10-01,Food,99\n".encode()
        self.post(SimpleUploadedFile("expenses.csv", csv_bytes))

        self.assertEqual(Expense.objects.get(user=self.user).amount, Decimal("99.00"))

    def test_any_invalid_row_saves_nothing_and_reports_rows(self):
        upload = _xlsx_upload([
            HEADERS,
            _row("2026-10-01", ["Food"], 10),
            _row("not a date", ["Food"], 10),
            _row("2026-10-01", ["Fodo"], 10),
            _row("2026-10-01", ["Food"], -5),
            _row("2026-10-01", ["Food"], 10, "Cheque"),
            _row("2026-10-01", ["Food", "Dining Out", "Tea"], 10),
        ])
        response = self.post(upload)

        self.assertEqual(response.status_code, 200)
        self.assertFalse(Expense.objects.exists())
        self.assertEqual([e.row for e in response.context["row_errors"]], [3, 4, 5, 6, 7])

    def test_subcategory_columns_must_be_filled_left_to_right(self):
        response = self.post(_xlsx_upload([HEADERS, _row("2026-10-01", ["Food", None, "Coffee"], 10)]))

        self.assertIn("Subcategory 1 is empty", response.context["row_errors"][0].message)

    def test_category_column_is_required(self):
        response = self.post(_xlsx_upload([HEADERS, _row("2026-10-01", [None, "Dining Out"], 10)]))

        self.assertIn("Category is required", response.context["row_errors"][0].message)

    def test_create_missing_categories_from_columns(self):
        upload = _xlsx_upload([
            HEADERS,
            _row("2026-10-01", ["Travel", "Taxi"], 300),
            _row("2026-10-02", ["Travel", "Taxi"], 200),
            _row("2026-10-03", ["Food", "Dining Out", "Tea"], 50),  # new leaf under existing parents
        ])
        self.post(upload, create_missing="1")

        taxi = Category.objects.get(user=self.user, name="Taxi")
        self.assertEqual(taxi.parent.name, "Travel")
        self.assertEqual(Expense.objects.filter(category=taxi).count(), 2)
        self.assertEqual(Category.objects.get(user=self.user, name="Tea").parent, self.dining)
        self.assertEqual(Category.objects.filter(user=self.user, name="Travel").count(), 1)

    def test_created_categories_roll_back_when_a_row_fails(self):
        upload = _xlsx_upload([HEADERS, _row("2026-10-01", ["Travel"], 300), _row("2026-10-02", ["Food"], "abc")])
        self.post(upload, create_missing="1")

        self.assertFalse(Category.objects.filter(name="Travel").exists())
        self.assertFalse(Expense.objects.exists())

    def test_ambiguous_bare_subcategory_name_is_an_error(self):
        travel = Category.objects.create(user=self.user, name="Travel")
        Category.objects.create(user=self.user, name="Coffee", parent=travel)  # a second "Coffee"
        response = self.post(_xlsx_upload([HEADERS, _row("2026-10-01", ["Coffee"], 10)]))

        self.assertIn("several subcategories", response.context["row_errors"][0].message)

    def test_root_category_wins_over_same_named_subcategory(self):
        root_coffee = Category.objects.create(user=self.user, name="Coffee")
        self.post(_xlsx_upload([HEADERS, _row("2026-10-01", ["Coffee"], 10)]))

        self.assertEqual(Expense.objects.get().category, root_coffee)

    def test_never_uses_another_users_categories(self):
        other = get_user_model().objects.create_user(email="o@example.com", password="pw")
        Category.objects.create(user=other, name="Rent")
        response = self.post(_xlsx_upload([HEADERS, _row("2026-10-01", ["Rent"], 10)]))

        self.assertEqual(len(response.context["row_errors"]), 1)
        self.assertFalse(Expense.objects.exists())

    def test_rejects_missing_headers_and_wrong_file_type(self):
        response = self.post(_xlsx_upload([["Amount"], [10]]))
        self.assertIn("missing: Date, Category", response.context["file_error"])

        response = self.post(SimpleUploadedFile("expenses.pdf", b"%PDF"))
        self.assertIn(".xlsx", response.context["file_error"])

    def test_template_has_columns_and_lists_categories(self):
        from openpyxl import load_workbook

        response = self.client.get(reverse("expenses:import_template"))
        workbook = load_workbook(BytesIO(response.content))

        self.assertEqual([c.value for c in workbook["Expenses"][1]], HEADERS)
        self.assertEqual(
            [[v for v in r if v] for r in workbook["Your categories"].iter_rows(min_row=2, values_only=True)],
            [["Food"], ["Food", "Dining Out"], ["Food", "Dining Out", "Coffee"]],
        )
        # The payment-type dropdown sits on the Payment Type column (I).
        self.assertTrue(str(workbook["Expenses"].data_validations.dataValidation[0].sqref).startswith("I2:I"))
